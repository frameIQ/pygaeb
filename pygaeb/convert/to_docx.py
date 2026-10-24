"""Word export: an LV as an editable .docx in the layout of the PDF printout.

Uses the same walk as the PDF printout: the same six columns, headings, sums,
"Nur Einh.-Preis" rule and markup rates. Word paginates the document itself, so
the header row repeats on every page but there are no carry-forward lines.

Requires python-docx: ``pip install pyGAEB[docx]``.

Usage::

    from pygaeb import GAEBParser
    from pygaeb.convert.to_docx import to_docx

    doc = GAEBParser.parse("tender.X83")
    to_docx(doc, "tender.docx")
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Literal

from pygaeb.convert._lv import Heading, Lv, Position, Total, build_lv
from pygaeb.models.document import GAEBDocument

_TEXT_WIDTH_CM = 18.0  # A4 minus 1.5 cm margins
_COLUMNS = (2.6, 8.8, 2.0, 1.2, 1.8, 2.2)  # same proportions as the PDF printout
_NUMERIC = (2, 4, 5)


def to_docx(
    doc: GAEBDocument,
    path: str | Path,
    title: str | None = None,
    *,
    locale: Literal["de", "en"] = "de",
    long_text: bool = True,
) -> None:
    """Render a GAEB document as a Word document.

    Args:
        doc: The document to render.
        path: Output .docx file path.
        title: Document title (defaults to the project name).
        locale: Labels and number format: ``"de"`` (default) or ``"en"``.
        long_text: Include the long texts.

    Raises:
        ImportError: If ``python-docx`` is not installed.
    """
    try:
        import docx
    except ImportError as e:
        raise ImportError(
            "Word export requires python-docx. Install with: pip install pyGAEB[docx]"
        ) from e

    lv = build_lv(doc, locale, title)
    document = docx.Document()
    _Writer(document, lv, long_text).write()
    document.save(str(Path(path)))


class _Writer:
    def __init__(self, document: Any, lv: Lv, long_text: bool) -> None:
        from docx.shared import Cm

        self.document = document
        self.lv = lv
        self.long_text = long_text
        scale = _TEXT_WIDTH_CM / sum(_COLUMNS)
        self.widths = [Cm(width * scale) for width in _COLUMNS]

    def write(self) -> None:
        from docx.shared import Cm, Pt

        section = self.document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(1.5))
        normal = self.document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(9)
        normal.paragraph_format.space_after = Pt(1)
        self.document.core_properties.title = self.lv.title

        title = self.document.add_paragraph()
        run = title.add_run(self.lv.title)
        run.bold = True
        run.font.size = Pt(14)
        for line in self.lv.meta:
            self.document.add_paragraph(line)

        self.table = self.document.add_table(rows=0, cols=len(_COLUMNS))
        self.table.autofit = False
        # Word reads the cell widths, LibreOffice the grid's: set both.
        for column, width in zip(self.table.columns, self.widths, strict=True):
            column.width = width
        self._header()
        for block in self.lv.blocks:
            if isinstance(block, Heading):
                self._heading(block)
            elif isinstance(block, Total):
                self._total(block)
            else:
                self._position(block)

    def _row(self) -> list[Any]:
        cells = list(self.table.add_row().cells)
        for cell, width in zip(cells, self.widths, strict=True):
            cell.width = width
        return cells

    def _header(self) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        labels = self.lv.labels
        cells = self._row()
        texts = (labels["oz"], labels["text"], labels["qty"], labels["unit"],
                 labels["unit_price"], labels["total"])
        for index, (cell, text) in enumerate(zip(cells, texts, strict=True)):
            _put(cell, text, bold=True, right=index in _NUMERIC)
            _border(cell, "bottom")
        header = OxmlElement("w:tblHeader")  # repeat on every page
        header.set(qn("w:val"), "true")
        self.table.rows[-1]._tr.get_or_add_trPr().append(header)

    def _heading(self, block: Heading) -> None:
        from docx.shared import Pt

        cells = self._row()
        _put(cells[0], block.number, bold=True)
        label = cells[1].merge(cells[-1])
        _put(label, block.label, bold=True)
        for cell in (cells[0], label):
            cell.paragraphs[0].paragraph_format.space_before = Pt(6)

    def _position(self, block: Position) -> None:
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

        cells = self._row()
        _put(cells[0], block.number)
        _put(cells[1], block.short_text, bold=True)
        if self.long_text:
            for line in block.long_text.splitlines():
                if line.strip():
                    _add(cells[1], line.strip())
        if block.marker:
            _add(cells[1], block.marker, italic=True)
        values = (block.qty, block.unit, block.unit_price, block.total)
        for index, value in enumerate(values, start=2):
            _put(cells[index], value or "", right=index in _NUMERIC)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    def _total(self, block: Total) -> None:
        cells = self._row()
        label = cells[1].merge(cells[4])
        _put(label, block.label, bold=True)
        _put(cells[5], block.amount, bold=True, right=True)
        if block.kind == "grand_total":
            for cell in (cells[0], label, cells[5]):
                _border(cell, "top")


def _put(cell: Any, text: str, *, bold: bool = False, italic: bool = False,
         right: bool = False) -> None:
    """Write ``text`` into the cell's first paragraph."""
    _format(cell.paragraphs[0], text, bold=bold, italic=italic, right=right)


def _add(cell: Any, text: str, *, italic: bool = False) -> None:
    _format(cell.add_paragraph(), text, italic=italic)


def _format(paragraph: Any, text: str, *, bold: bool = False, italic: bool = False,
            right: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if text:
        run = paragraph.add_run(text)
        run.bold = bold or None
        run.italic = italic or None
    if right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _border(cell: Any, edge: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        # tcPr children have a fixed order; tcBorders comes before shading and alignment.
        tc_pr.insert_element_before(
            borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
            "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
            "w:tcPrChange",
        )
    line = OxmlElement(f"w:{edge}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), "4")
    line.set(qn("w:color"), "000000")
    borders.append(line)


__all__ = ["to_docx"]
